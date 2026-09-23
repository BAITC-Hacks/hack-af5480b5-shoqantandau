import shutil

import numpy as np
import pandas as pd
from django.contrib import messages
from django.http import Http404
from django.shortcuts import redirect, render
from django.views.decorators.http import require_POST

import constants
from django.contrib.auth.views import LoginView
from django.core.exceptions import PermissionDenied

from . import loaders


def _need(request, perm: str):
    """Проверка права роли; без права — понятная страница «Недостаточно прав»."""
    if not request.user.has_perm(f"procurement.{perm}"):
        raise PermissionDenied(perm)


class DemoLoginView(LoginView):
    template_name = "procurement/login.html"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["demo_users"] = [
            ("manager", "manager12345", "Менеджер закупа", "считает, правит, утверждает и скачивает заказ"),
            ("viewer", "viewer12345", "Наблюдатель", "только просмотр: руководитель, склад"),
            ("admin", "admin12345", "Администратор", "всё + загрузка данных 1С и пользователи"),
        ] if constants.SHOW_DEMO_USERS else []
        return ctx


def _data_status():
    """Какие файлы данных лежат на месте — по каждому поставщику."""
    result = []
    for key, sup in constants.SUPPLIERS.items():
        files = []
        resolved = loaders.resolve_files(key)
        for kind, fname in constants.DATA_FILES.items():
            path = resolved[kind]
            files.append({"kind": kind, "name": fname, "exists": path.exists(),
                          "uploaded": loaders.UPLOAD_DIR in path.parents,
                          "size_kb": round(path.stat().st_size / 1024) if path.exists() else 0})
        result.append({"key": key, "name": sup["name"], "files": files,
                       "ready": all(f["exists"] for f in files)})
    return result


def index(request):
    suppliers = _data_status()
    for s in suppliers:
        if not s["ready"]:
            continue
        try:
            data = loaders.get_supplier(s["key"])
            s["summary"] = data.summary()
            s["warnings"] = data.warnings
        except Exception as exc:  # показываем ошибку разбора файла, а не падаем
            s["error"] = f"{type(exc).__name__}: {exc}"
    kinds = {"sales_transactions": "Динамика продаж (накладные)", "sales_monthly": "Ежемесячные продажи",
             "stock_monthly": "Ежемесячные остатки", "in_transit": "Товар в пути", "moq": "MOQ / кратность",
             "seasonality": "Сезонность"}
    return render(request, "procurement/index.html", {"suppliers": suppliers, "kinds": kinds,
                                                      "has_uploads": loaders.UPLOAD_DIR.exists()})


@require_POST
def upload_data(request):
    """Загрузка новой выгрузки 1С вместо демо-файла. Если файл не разбирается — откат."""
    _need(request, "manage_data")
    sup = request.POST.get("supplier")
    kind = request.POST.get("kind")
    f = request.FILES.get("file")
    if sup not in constants.SUPPLIERS or kind not in constants.DATA_FILES or not f:
        messages.error(request, "Выберите поставщика, тип выгрузки и xlsx-файл.")
        return redirect("procurement:index")
    if not f.name.lower().endswith(".xlsx"):
        messages.error(request, "Нужен файл .xlsx — выгрузка из 1С в Excel.")
        return redirect("procurement:index")
    target = loaders.UPLOAD_DIR / sup / constants.DATA_FILES[kind]
    target.parent.mkdir(parents=True, exist_ok=True)
    backup = target.read_bytes() if target.exists() else None
    with open(target, "wb") as out:
        for chunk in f.chunks():
            out.write(chunk)
    try:
        loaders.get_supplier(sup, use_cache=False)
    except Exception as exc:
        if backup is None:
            target.unlink(missing_ok=True)
        else:
            target.write_bytes(backup)
        messages.error(request, f"Файл не подходит по формату, оставлены прежние данные. Ошибка: {exc}")
        return redirect("procurement:index")
    messages.success(request, f"{constants.SUPPLIERS[sup]['name']}: загружен «{f.name}» как {constants.DATA_FILES[kind]}. "
                              f"Следующий расчёт использует новые данные.")
    return redirect("procurement:index")


@require_POST
def reset_uploads(request):
    _need(request, "manage_data")
    shutil.rmtree(loaders.UPLOAD_DIR, ignore_errors=True)
    shutil.rmtree(loaders.CACHE_DIR, ignore_errors=True)
    loaders._MEMORY.clear()
    messages.info(request, "Загруженные файлы удалены, используются демо-данные партнёра.")
    return redirect("procurement:index")


@require_POST
def reload_data(request):
    _need(request, "manage_data")
    shutil.rmtree(loaders.CACHE_DIR, ignore_errors=True)
    loaders._MEMORY.clear()
    messages.info(request, "Данные перечитаны из файлов.")
    return redirect("procurement:index")


def sku_detail(request, supplier, code):
    if supplier not in constants.SUPPLIERS:
        raise Http404("Неизвестный поставщик")
    data = loaders.get_supplier(supplier)
    if code not in data.items.index:
        raise Http404("Артикул не найден")
    item = data.items.loc[code]

    tx = data.transactions[data.transactions["code"] == code]
    tx_m = tx.groupby(tx["date"].dt.to_period("M"))["qty"].sum()
    m = data.monthly[data.monthly["code"] == code].set_index("month").sort_index()
    rows = [{"month": p.strftime("%m.%Y"), "sales_1c": r.sales_1c, "tx": float(tx_m.get(p, 0)),
             "stock_open": r.stock_open} for p, r in m.iterrows()]

    top_docs = (tx.groupby(["doc", tx["date"].dt.date])["qty"].sum()
                .sort_values(ascending=False).head(10).reset_index())
    top_docs.columns = ["doc", "date", "qty"]

    # график: сырые продажи против очищенного спроса (тот же расчёт, что и в заказе)
    from . import charts, engine as eng
    chart, calc = "", None
    res = eng.calculate(data, eng.Params(codes=[code]))
    if not res.empty:
        calc = res.iloc[0]
        det = calc["details"]
        months = det["months"]
        so = {months.index(f"{y}-{m:02d}") for y, m in
              [(int(x.split()[1]), eng.MONTH_NAMES.index(x.split()[0]) + 1) for x in det["stockout_months"]]
              if f"{y}-{m:02d}" in months}
        oo = {months.index(o["month"]) for o in det["one_offs"] if o["month"] in months}
        chart = charts.demand_chart(months, det["raw"], det["clean"], so, oo)
    return render(request, "procurement/sku.html", {
        "chart": chart, "calc": calc,
        "supplier": data.name, "supplier_key": supplier, "code": code, "item": item,
        "rows": rows, "top_docs": top_docs.to_dict("records"),
        "next_arrival": None if pd.isna(item["next_arrival"]) else pd.Timestamp(item["next_arrival"]).date(),
    })


def catalog(request):
    """Каталог товаров обоих поставщиков: поиск, фильтры и сортировка прямо в браузере."""
    rows, groups = [], {}
    for key in constants.SUPPLIERS:
        data = loaders.get_supplier(key)
        last_full = pd.Period(data.data_date, freq="M") - 1
        m = data.monthly
        s12 = m[(m["month"] > last_full - 12) & (m["month"] <= last_full)].groupby("code")["sales_1c"].sum()
        s1 = m[m["month"] == last_full].groupby("code")["sales_1c"].sum()
        it = data.items
        for code, r in it.iterrows():
            g = str(r["category"])
            groups.setdefault(g, r["name"][:32])
            rows.append({"sup": key, "supName": data.name, "code": code, "article": str(r["article"] or ""),
                         "name": r["name"], "group": g, "unit": r["unit"],
                         "stock": round(float(r["stock_now"] or 0)), "transit": round(float(r["in_transit"] or 0)),
                         "moq": int(r["moq"] or 1), "s12": round(float(s12.get(code, 0))), "s1": round(float(s1.get(code, 0))),
                         "marked": bool(r.get("marked", False)),
                         "cost": round(float(r["cost"]), 2) if pd.notna(r.get("cost")) else None})
    return render(request, "procurement/catalog.html", {
        "rows": rows, "suppliers": constants.SUPPLIERS,
        "groups": sorted(groups.items()), "q": request.GET.get("q", ""),
    })


def sku_search(request):
    q = request.GET.get("q", "").strip()
    results = []
    if q:
        for key in constants.SUPPLIERS:
            items = loaders.get_supplier(key).items
            mask = (items.index.str.contains(q, case=False, regex=False)
                    | items["name"].str.contains(q, case=False, regex=False)
                    | items["article"].astype(str).str.contains(q, case=False, regex=False))
            for code, r in items[mask].head(50).iterrows():
                results.append({"supplier_key": key, "supplier": constants.SUPPLIERS[key]["name"],
                                "code": code, "name": r["name"], "article": r["article"]})
    return render(request, "procurement/search.html", {"q": q, "results": results})


# ---------- расчёт ----------

from django.core.paginator import Paginator  # noqa: E402
from django.shortcuts import get_object_or_404  # noqa: E402

from . import engine, services  # noqa: E402
from .models import CalculationRun  # noqa: E402


def _num(v, cast=float):
    try:
        return cast(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def run_form(request):
    if request.method == "POST":
        _need(request, "run_calculation")
        f = request.POST
        lead_times = {k: v for k in constants.SUPPLIERS if (v := _num(f.get(f"lead_{k}"), int))}
        params = engine.Params(
            lead_times=lead_times,
            review_days=_num(f.get("review_days"), int) or constants.REVIEW_PERIOD_DAYS,
            growth_plan_pct=_num(f.get("growth_plan_pct")),
            category=f.get("category", "").strip(),
            use_outliers=bool(f.get("use_outliers")),
            use_stockout=bool(f.get("use_stockout")),
            use_seasonality=bool(f.get("use_seasonality")),
            use_growth=bool(f.get("use_growth")),
        )
        sup = f.get("supplier") or ""
        run = services.run_calculation(params, [sup] if sup in constants.SUPPLIERS else None, user=request.user)
        return redirect("procurement:run_detail", run.pk)
    # группы товаров для выбора — только если данные уже разобраны (иначе страница открывалась бы 20–30 с)
    groups = []
    cached = loaders.is_ready() or all((loaders.CACHE_DIR / f"{k}.pkl").exists() for k in constants.SUPPLIERS)
    leads = {k: v.get("lead_time_days") for k, v in constants.SUPPLIERS.items()}
    if cached:
        seen = {}
        for k in constants.SUPPLIERS:
            data = loaders.get_supplier(k)
            leads[k] = data.lead_time_days
            for cat, g in data.items.groupby("category"):
                e = seen.setdefault(cat, {"code": cat, "n": 0, "example": g["name"].iloc[0][:40], "sup": set()})
                e["n"] += len(g)
                e["sup"].add(data.name)
        groups = sorted(seen.values(), key=lambda e: -e["n"])
        for e in groups:
            e["sup"] = ", ".join(sorted(e["sup"]))
    return render(request, "procurement/run_form.html", {
        "suppliers": [(k, v["name"], leads[k]) for k, v in constants.SUPPLIERS.items()],
        "review_days": constants.REVIEW_PERIOD_DAYS, "groups": groups, "cached": cached,
        "runs": CalculationRun.objects.select_related("created_by")[:15],
    })


def _apply_edits(request, run):
    """Сохранить правки количества и/или изменить статус выбранных строк."""
    action = request.POST.get("action", "save")
    ids = [int(x) for x in request.POST.getlist("sel") if x.isdigit()]
    changed = 0
    for key, val in request.POST.items():
        if not key.startswith("qty_"):
            continue
        line_id = _num(key[4:], int)
        new = _num(val, int)
        if line_id is None:
            continue
        line = run.lines.filter(pk=line_id).first()
        if line is None:
            continue
        final = None if new is None or new == line.qty_recommended else max(new, 0)
        if final != line.qty_final:
            line.qty_final = final
            line.save(update_fields=["qty_final"])
            changed += 1
    if action in ("approve", "reject", "reset") and ids:
        status = {"approve": "approved", "reject": "rejected", "reset": "new"}[action]
        from django.utils import timezone
        n = run.lines.filter(pk__in=ids).update(
            status=status, decided_by=None if action == "reset" else request.user,
            decided_at=None if action == "reset" else timezone.now())
        label = {"approve": "Утверждено товаров", "reject": "Отмечено «не заказываем»",
                 "reset": "Возвращено в работу"}[action]
        messages.success(request, f"{label}: {n}.")
    elif action in ("approve", "reject", "reset"):
        messages.error(request, "Сначала отметьте товары галочками слева, затем нажмите кнопку.")
    elif changed:
        messages.success(request, f"Изменённые количества сохранены: {changed}.")


VIEWS = [("order", "К заказу"), ("urgent", "Срочные"), ("review", "Проверить вручную"),
         ("approved", "Утверждённые"), ("all", "Весь список")]


def run_detail(request, pk):
    run = get_object_or_404(CalculationRun, pk=pk)
    if request.method == "POST":
        _need(request, "edit_order")
        _apply_edits(request, run)
        return redirect(request.get_full_path())
    sup_keys = [k for k in run.supplier.split(",") if k]
    active = request.GET.get("supplier") or (sup_keys[0] if sup_keys else "")
    view = request.GET.get("view", "order")
    stat = run.stats.get(active, {})
    fields = ["pk", "code_1c", "article", "name", "category", "stock", "in_transit", "regular_demand", "moq",
              "qty_recommended", "qty_final", "urgency", "days_of_cover", "status", "needs_review", "cost"]
    rows = []
    for l in run.lines.filter(supplier=active).only(*fields):
        final = l.qty_to_order
        rows.append({
            "id": l.pk, "code": l.code_1c, "article": l.article, "name": l.name, "group": l.category,
            "stock": round(l.stock), "transit": round(l.in_transit), "demand": round(l.regular_demand, 1),
            "moq": int(l.moq), "rec": l.qty_recommended, "qty": final, "edited": l.qty_final is not None,
            "urg": l.urgency, "urgRank": {"high": 0, "medium": 1, "low": 2}.get(l.urgency, 3),
            "cover": None if (l.days_of_cover or 0) >= 9999 else round(l.days_of_cover, 1),
            "status": l.status, "review": l.needs_review,
            "sum": round(l.cost * final) if l.cost else None,
        })
    rows.sort(key=lambda r: (r["urgRank"], r["cover"] if r["cover"] is not None else 1e9))
    approved = sum(1 for r in rows if r["status"] == "approved")
    return render(request, "procurement/run_detail.html", {
        "run": run, "active": active, "tabs": [(k, run.stats.get(k, {})) for k in sup_keys],
        "stat": stat, "view": view, "views": VIEWS, "rows": rows, "approved": approved,
        "value_saved": (stat.get("value_raw", 0) or 0) - (stat.get("value_total", 0) or 0),
        "qty_saved": (stat.get("qty_raw_total", 0) or 0) - (stat.get("qty_total", 0) or 0),
        "can_edit": request.user.has_perm("procurement.edit_order"),
        "can_export": request.user.has_perm("procurement.export_order"),
    })


def line_reason(request, pk):
    """Обоснование строки — подгружается при раскрытии, чтобы список открывался мгновенно."""
    from .templatetags.procurement_extras import sentences
    l = get_object_or_404(OrderLine.objects.select_related("decided_by"), pk=pk)
    who = ""
    if l.decided_by and l.decided_at:
        from django.utils import timezone
        who = (f"{'Утвердил' if l.status == 'approved' else 'Отметил «не заказываем»'}: "
               f"{l.decided_by.get_full_name() or l.decided_by.username}, "
               f"{timezone.localtime(l.decided_at):%d.%m.%Y %H:%M}")
    return JsonResponse({"sentences": sentences(l.reason), "ai": l.reason_ai, "who": who,
                         "supplier": l.supplier, "code": l.code_1c})


# ---------- экспорт ----------

import io  # noqa: E402

from django.http import HttpResponse  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402


def export_run(request, pk):
    """Выгрузка в xlsx: лист на каждого поставщика, по умолчанию только утверждённые позиции."""
    _need(request, "export_order")
    run = get_object_or_404(CalculationRun, pk=pk)
    scope = request.GET.get("scope", "approved")
    only = request.GET.get("supplier")
    wb = Workbook()
    wb.remove(wb.active)
    headers = ["Код 1С", "Артикул поставщика", "Наименование", "Количество", "Ед.", "Кратность",
               "Срочность", "Остаток", "В пути", "Рекомендовано системой", "Себестоимость, ₸", "Сумма, ₸", "Обоснование"]
    total = 0
    for key in [k for k in run.supplier.split(",") if k and (not only or k == only)]:
        qs = run.lines.filter(supplier=key)
        qs = qs.filter(status="approved") if scope == "approved" else qs.filter(qty_recommended__gt=0).exclude(status="rejected")
        ws = wb.create_sheet(constants.SUPPLIERS.get(key, {}).get("name", key)[:31])
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        for l in qs.order_by("code_1c"):
            if l.qty_to_order <= 0:
                continue
            ws.append([l.code_1c, l.article, l.name, l.qty_to_order, (l.details or {}).get("unit", "шт"),
                       l.moq, l.get_urgency_display(), round(l.stock), round(l.in_transit),
                       l.qty_recommended, l.cost, round(l.value_to_order, 2) if l.value_to_order else None, l.reason])
            total += 1
        for col, w in zip("ABCDEFGHIJKLM", [14, 22, 60, 12, 6, 10, 11, 10, 10, 14, 14, 14, 100]):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
    if total == 0 and scope == "approved":
        messages.error(request, "Пока нет утверждённых товаров. Отметьте товары галочками и нажмите «Утвердить» "
                                "внизу экрана — или скачайте все рекомендации.")
        return redirect("procurement:run_detail", pk)
    buf = io.BytesIO()
    wb.save(buf)
    name = f"zakaz_{pk}_{'utverzhden' if scope == 'approved' else 'chernovik'}.xlsx"
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{name}"'
    return resp


# ---------- проверка для экспертов ----------

DEMO_CODES = {"iek": "030201020_", "se": "300200294_"}  # артикулы для демонстрации по умолчанию


def check_scenarios(request):
    """Один артикул в нескольких сценариях: видно, как каждый источник данных влияет на результат."""
    g = request.GET
    sup = g.get("supplier", "iek") if g.get("supplier") in constants.SUPPLIERS else "iek"
    code = (g.get("code") or DEMO_CODES.get(sup, "")).strip()
    data = loaders.get_supplier(sup)
    ctx = {"suppliers": constants.SUPPLIERS, "sup": sup, "code": code, "g": g}
    if code not in data.items.index:
        ctx["error"] = f"Артикул {code} не найден у поставщика {data.name}."
        return render(request, "procurement/check.html", ctx)

    base_p = dict(codes=[code])
    base = engine.calculate(data, engine.Params(**base_p))
    if base.empty:
        ctx["error"] = "У артикула нет продаж за последние 12 месяцев — заказ не рассчитывается."
        return render(request, "procurement/check.html", ctx)
    b = base.iloc[0]
    item = data.items.loc[code]
    big_qty = _num(g.get("big_qty")) or max(round(b["regular_demand"] * 10), 50)
    transit_add = _num(g.get("transit_add")) or max(round(b["qty_recommended"] / 2), 10)
    stock_val = _num(g.get("stock_val"))
    stock_val = stock_val if stock_val is not None else float(item["stock_now"]) + transit_add
    growth = _num(g.get("growth")) if g.get("growth") not in (None, "") else 30.0
    last_month = (pd.Period(data.data_date, freq="M") - 2).start_time + pd.Timedelta(days=10)
    test_order = [{"code": code, "qty": big_qty, "date": str(last_month.date()), "doc": "ТЕСТ-РАЗОВЫЙ"}]
    test_month = [{**test_order[0], "doc": "ТЕСТ-МЕСЯЦ", "monthly_only": True}]
    # группа с самым непохожим сезонным профилем — чтобы показать влияние категории товара
    profiles = engine.category_profiles(data, pd.Period(data.data_date, freq="M") - 1)
    my_cat = str(item["category"])
    other_cat = None
    if my_cat in profiles and len(profiles) > 1:
        other_cat = max((c for c in profiles if c != my_cat),
                        key=lambda c: float(np.abs(profiles[c][0] - profiles[my_cat][0]).sum()))

    scen = [
        ("base", "Базовый расчёт", "Все данные как есть", {}),
        ("oneoff", f"+ разовая продажа {big_qty:.0f} шт.", "Требование 4: крупная сделка не раздувает заказ",
         {"test_orders": test_order}),
        ("oneoff_raw", "То же, но без очистки данных", "Для сравнения: так посчитал бы Excel по сырым продажам",
         {"test_orders": test_order, "use_outliers": False}),
        ("spike", f"Всплеск {big_qty:.0f} шт. только в отчёте 1С", "Требование 4: всплеск без накладной тоже сглаживается",
         {"test_orders": test_month}),
        ("transit", f"+ {transit_add:.0f} шт. в пути", "Требование 1: товар в пути уменьшает заказ",
         {"overrides": {code: {"in_transit": float(item["in_transit"]) + transit_add}}}),
        ("stock", f"Остаток {stock_val:.0f} шт.", "Требование 1: остаток уменьшает заказ",
         {"overrides": {code: {"stock_now": stock_val}}}),
        ("growth", f"Рост продаж {growth:+.0f}% в год", "Требование 1: ожидаемый рост увеличивает заказ", {"growth_plan_pct": growth}),
    ] + ([
        ("category", f"Группа товаров {my_cat} → {other_cat}", "Требование 1: у другой группы товаров другой сезон и рост",
         {"overrides": {code: {"category": other_cat}}}),
    ] if other_cat else []) + [
        ("no_stockout", "Без учёта дефицита", "Требование 3: если не учитывать, что товара не было, заказ меньше", {"use_stockout": False}),
        ("no_season", "Без сезонности", "Требование 2: если считать по среднему, без учёта сезона", {"use_seasonality": False}),
    ]
    rows = []
    for key, title, hint, extra in scen:
        r = b if key == "base" else engine.calculate(data, engine.Params(**base_p, **extra)).iloc[0]
        req, _, plain = hint.partition(": ") if hint.startswith("Требование") else ("", "", hint)
        rows.append({"key": key, "title": title, "hint": plain, "req": req, "qty": int(r["qty_recommended"]),
                     "delta": int(r["qty_recommended"]) - int(b["qty_recommended"]),
                     "demand": r["regular_demand"], "forecast": r["forecast"], "reason": r["reason"]})
    ctx.update({"rows": rows, "item": item, "base": b, "big_qty": big_qty, "transit_add": transit_add,
                "stock_val": stock_val, "growth": growth, "supplier_name": data.name})
    return render(request, "procurement/check.html", ctx)


# ---------- объяснение простыми словами ----------

from django.http import JsonResponse  # noqa: E402

from . import llm  # noqa: E402
from .models import OrderLine  # noqa: E402


@require_POST
def explain_line(request, pk):
    line = get_object_or_404(OrderLine, pk=pk)
    text, source = llm.explain(line)
    line.reason_ai = text
    line.save(update_fields=["reason_ai"])
    return JsonResponse({"text": text, "source": source})


# ---------- бэктест ----------

from . import backtest  # noqa: E402


def backtest_view(request):
    if request.method == "POST":
        _need(request, "run_calculation")
        backtest.run()
        messages.success(request, "Бэктест пересчитан.")
        return redirect("procurement:backtest")
    res = backtest.load()
    rows = []
    if res:
        for key, v in res["suppliers"].items():
            sv, ex = v["summary"].get("service", {}), v["summary"].get("excel", {})
            rows.append({"name": v["name"], "service": sv, "excel": ex, "per_month": v["per_month"],
                         "over_cut": (1 - sv["over"] / ex["over"]) if ex.get("over") else None,
                         "wape_gain": (ex["wape"] - sv["wape"]) if ex.get("wape") is not None else None})
    return render(request, "procurement/backtest.html", {"res": res, "rows": rows})
