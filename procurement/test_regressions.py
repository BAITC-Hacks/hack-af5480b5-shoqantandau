"""Failures that the original happy-path tests did not cover."""
from datetime import date
from io import BytesIO
from unittest.mock import patch

import numpy as np
import pandas as pd
from django.contrib.auth.models import User
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook

from .engine import Params, calculate, first_shortage, prepare_lines, detect_one_offs_fast
from .models import CalculationRun, OrderLine
from .tests import make_data, row
from .loaders import apply_stockout_periods, read_current_stock


class OptionalSources(SimpleTestCase):
    def test_overlapping_periods_count_days_once(self):
        raw = pd.DataFrame([
            ["Код 1с", "Начало", "Конец"],
            ["A", pd.Timestamp("2026-06-01"), pd.Timestamp("2026-06-10")],
            ["A", pd.Timestamp("2026-06-05"), pd.Timestamp("2026-06-15")],
        ])
        data = make_data({"A": {"monthly": [30.] * 33}})
        with patch("procurement.loaders._read_raw", return_value=raw):
            monthly = apply_stockout_periods(data.monthly, None, data.data_date)
        june = monthly.loc[monthly.month == pd.Period("2026-06"), "stockout_days"].iloc[0]
        self.assertEqual(june, 15)

    def test_exact_stockout_overrides_positive_monthly_stock(self):
        data = make_data({"A": {"monthly": [30.] * 33}})
        before = row(calculate(data), "A")
        data.monthly["stockout_days"] = 0.
        data.monthly.loc[data.monthly.month == pd.Period("2026-06"), "stockout_days"] = 15
        after = row(calculate(data), "A")
        self.assertGreater(after.stockout_added, before.stockout_added)
        self.assertGreater(after.forecast, before.forecast)

    def test_current_stock_requires_matching_date(self):
        raw = pd.DataFrame([["Код 1с", "Свободный остаток", "Дата"], ["A", 45, pd.Timestamp("2026-09-22")]])
        with patch("procurement.loaders._read_raw", return_value=raw):
            self.assertEqual(read_current_stock(None, date(2026, 9, 22))["A"], 45)
            with self.assertRaises(ValueError):
                read_current_stock(None, date(2026, 9, 23))

    def test_duplicate_stock_keys_are_rejected(self):
        raw = pd.DataFrame([["Код 1с", "Свободный остаток", "Дата"],
                            ["A", 45, pd.Timestamp("2026-09-22")], ["A", 50, pd.Timestamp("2026-09-22")]])
        with patch("procurement.loaders._read_raw", return_value=raw), self.assertRaises(ValueError):
            read_current_stock(None, date(2026, 9, 22))

    def test_invalid_stockout_interval_rejected(self):
        raw = pd.DataFrame([["Код 1с", "Начало", "Конец"],
                            ["A", pd.Timestamp("2026-06-15"), pd.Timestamp("2026-06-01")]])
        data = make_data({"A": {"monthly": [30.] * 33}})
        with patch("procurement.loaders._read_raw", return_value=raw), self.assertRaises(ValueError):
            apply_stockout_periods(data.monthly, None, data.data_date)


class SupplyDates(SimpleTestCase):
    def data(self):
        return make_data({"A": {"monthly": [30.0] * 33, "stock_now": 2, "in_transit": 200}})

    def test_late_transit_does_not_reduce_order(self):
        data = self.data()
        data.shipments = pd.DataFrame([{"code": "A", "qty": 200, "arrival_date": pd.Timestamp("2027-01-01")}])
        result = row(calculate(data, Params(lead_time_days=30)), "A")
        without = row(calculate(data, Params(lead_time_days=30, overrides={"A": {"in_transit": 0}})), "A")
        self.assertEqual(result.qty_recommended, without.qty_recommended)
        self.assertEqual(result.details["late_transit"], 200)

    def test_split_shipments_keep_individual_dates(self):
        data = self.data()
        data.shipments = pd.DataFrame([
            {"code": "A", "qty": 10, "arrival_date": pd.Timestamp("2026-09-23")},
            {"code": "A", "qty": 190, "arrival_date": pd.Timestamp("2027-01-01")},
        ])
        result = row(calculate(data, Params(lead_time_days=30)), "A")
        self.assertEqual(result.details["eligible_transit"], 10)
        self.assertEqual(result.details["late_transit"], 190)

    def test_gap_before_receipt_is_urgent_even_when_total_covers(self):
        data = self.data()
        data.shipments = pd.DataFrame([{"code": "A", "qty": 200, "arrival_date": pd.Timestamp("2026-10-10")}])
        result = row(calculate(data, Params(lead_time_days=30)), "A")
        self.assertEqual(result.qty_recommended, 0)
        self.assertEqual(result.urgency, "high")
        self.assertLess(result.details["first_shortage_day"], 3)

    def test_today_receipt_prevents_gap(self):
        self.assertIsNone(first_shortage(0, [(date(2026, 9, 22), 50)], 1,
                                        date(2026, 9, 22), 30, np.ones(12), 1))

    def test_overdue_shipment_is_not_counted_twice(self):
        data = self.data()
        data.shipments = pd.DataFrame([{"code": "A", "qty": 200, "arrival_date": pd.Timestamp("2026-09-01")}])
        result = row(calculate(data, Params(lead_time_days=30)), "A")
        self.assertEqual(result.details["eligible_transit"], 0)
        self.assertTrue(result.needs_review)


class ForecastBoundaries(SimpleTestCase):
    def test_partial_or_future_seasonality_cannot_change_forecast(self):
        data = make_data({"A": {"monthly": [30.] * 33}})
        p = Params(use_growth=False)
        before = row(calculate(data, p), "A").forecast
        data.seasonality.loc[2026, 9:12] = 1e12
        after = row(calculate(data, p), "A").forecast
        self.assertAlmostEqual(before, after)

    def test_minus_100_growth_means_zero_demand(self):
        data = make_data({"A": {"monthly": [30.] * 33, "stock_now": 0}})
        result = row(calculate(data, Params(growth_plan_pct=-100)), "A")
        self.assertEqual(result.forecast, 0)
        self.assertEqual(result.qty_recommended, 0)

    def test_invalid_parameters_rejected(self):
        for params in (Params(lead_time_days=-1), Params(growth_plan_pct=float("nan")),
                       Params(growth_plan_pct=float("inf")), Params(review_days=100000),
                       Params(overrides={"A": {"stock_now": -1}})):
            with self.subTest(params=params), self.assertRaises(ValueError):
                params.validate()

    def test_customer_split_invoices_are_aggregated(self):
        tx = pd.DataFrame([
            {"code": "A", "doc": f"regular-{i}", "date": pd.Timestamp(2026, 6, i + 1), "qty": 5, "client_id": f"anon-{i}"}
            for i in range(10)
        ] + [
            {"code": "A", "doc": f"big-{i}", "date": pd.Timestamp(2026, 6, 20, i), "qty": 100, "client_id": "anon-bulk"}
            for i in range(3)
        ])
        lines = prepare_lines(tx, [pd.Period("2026-06")])["A"]
        flagged = detect_one_offs_fast(lines, np.array([20.] * 24), 5, .4)
        self.assertEqual(len(flagged), 1)
        self.assertEqual(flagged[0]["qty"], 300)


class OrderWorkflow(TestCase):
    def setUp(self):
        self.client.force_login(User.objects.get(username="manager"))
        self.run = CalculationRun.objects.create(supplier="iek")
        self.line = OrderLine.objects.create(run=self.run, supplier="iek", code_1c="A", name="Тест",
                                             qty_recommended=20, moq=10, status="approved")

    def test_edit_resets_approval_and_ai(self):
        self.line.reason_ai = "Old quantity"
        self.line.save()
        self.client.post(f"/run/{self.run.pk}/", {"action": "save", f"qty_{self.line.pk}": "30"})
        self.line.refresh_from_db()
        self.assertEqual(self.line.status, "new")
        self.assertEqual(self.line.qty_final, 30)
        self.assertEqual(self.line.reason_ai, "")

    def test_bad_quantity_does_not_modify_or_approve(self):
        for value in ("abc", "nan", "-2", "2.5", "21", "9999999999999"):
            self.line.status = "new"
            self.line.save()
            self.client.post(f"/run/{self.run.pk}/", {
                "action": "approve", "sel": [self.line.pk], f"qty_{self.line.pk}": value})
            self.line.refresh_from_db()
            self.assertIsNone(self.line.qty_final)
            self.assertEqual(self.line.status, "new")

    def test_edit_and_explicit_approve_in_one_request(self):
        self.client.post(f"/run/{self.run.pk}/", {
            "action": "approve", "sel": [self.line.pk], f"qty_{self.line.pk}": "30"})
        self.line.refresh_from_db()
        self.assertEqual(self.line.status, "approved")
        self.assertEqual(self.line.qty_to_order, 30)
        self.assertIsNotNone(self.line.decided_by_id)

    def test_invalid_export_returns_404(self):
        self.assertEqual(self.client.get(f"/run/{self.run.pk}/export/?supplier=unknown").status_code, 404)
        self.assertEqual(self.client.get(f"/run/{self.run.pk}/export/?scope=unknown").status_code, 404)

    def test_exported_names_are_literal_text(self):
        self.line.name = '=HYPERLINK("https://example.org","unsafe")'
        self.line.save()
        response = self.client.get(f"/run/{self.run.pk}/export/")
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(wb.active["C2"].data_type, "s")
        self.assertEqual(wb.active["D2"].value, 20)

    def test_manual_order_is_included_in_draft(self):
        self.line.qty_recommended = 0
        self.line.qty_final = 30
        self.line.save()
        response = self.client.get(f"/run/{self.run.pk}/export/?scope=draft")
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb.active["D2"].value, 30)

    def test_invalid_calculation_does_not_create_run(self):
        for value in ("nan", "inf", "-101", "999999"):
            response = self.client.post("/", {"review_days": "30", "growth_plan_pct": value})
            self.assertEqual(response.status_code, 302)
        self.assertEqual(CalculationRun.objects.count(), 1)

    def test_viewer_cannot_invoke_external_ai(self):
        self.client.force_login(User.objects.get(username="viewer"))
        with patch("procurement.llm.explain") as explain:
            self.assertEqual(self.client.post(f"/line/{self.line.pk}/explain/").status_code, 403)
            explain.assert_not_called()

    def test_failed_calculation_rolls_back_run(self):
        from .services import run_calculation
        with patch("procurement.services.loaders.get_supplier", side_effect=FileNotFoundError("missing")):
            with self.assertRaises(FileNotFoundError):
                run_calculation(Params(), ["iek"])
        self.assertEqual(CalculationRun.objects.count(), 1)
